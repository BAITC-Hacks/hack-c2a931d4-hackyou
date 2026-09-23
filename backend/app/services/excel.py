"""An Excel presentation of verified CSV exports, separate from the engine snapshot."""

import csv
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.errors import AppError
from backend.app.services.analyses import AnalysisService

SHEETS = {
    "top_nodes.csv": "Приоритеты",
    "nodes_roles.csv": "Участники",
    "clusters.csv": "Сообщества",
}
NUMERIC = {"rank", "role_score", "priority_score", "n_nodes", "n_seed", "sum_kzt_internal"}
HEADERS = {
    "gid": "GID участника",
    "role": "Роль (код движка)",
    "role_score": "Оценка роли",
    "priority_score": "Приоритет",
    "cluster_id": "ID сообщества",
    "evidence": "Наблюдения",
    "rank": "Место",
    "why": "Обоснование",
    "n_nodes": "Участников",
    "n_seed": "Исходных узлов (seed)",
    "sum_kzt_internal": "Внутренние переводы, ₸",
    "top_gids": "Ключевые GID",
    "hypothesis": "Гипотеза",
}


def excel_report(analyses: AnalysisService, analysis_id: str) -> bytes:
    # Verify every source before creating a derivative. Canonical bytes/hashes stay untouched.
    paths = {name: analyses.download(analysis_id, name) for name in SHEETS}
    workbook = Workbook(write_only=True)
    workbook.properties.title = "TraceGraph — результаты анализа"
    workbook.properties.description = f"Снимок анализа {analysis_id}. Источники: три CSV движка."
    for filename, title in SHEETS.items():
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        with paths[filename].open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            fields = reader.fieldnames or []
            headers = []
            for index, field in enumerate(fields, 1):
                sheet.column_dimensions[get_column_letter(index)].width = (
                    85
                    if field in {"evidence", "why", "hypothesis"}
                    else 52
                    if field == "top_gids"
                    else 24
                    if field in {"gid", "role", "sum_kzt_internal"}
                    else 20
                )
                cell = WriteOnlyCell(sheet, HEADERS.get(field, field))
                cell.data_type = "s"
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="426337")
                headers.append(cell)
            sheet.append(headers)
            row_number = 1
            for row_number, row in enumerate(reader, 2):
                if row_number > 1_048_576:
                    raise AppError(
                        "excel_row_limit",
                        "Слишком много строк для листа Excel. Используйте CSV.",
                        422,
                    )
                cells = []
                for field in fields:
                    value = row[field]
                    if len(value) > 32_767:
                        raise AppError(
                            "excel_cell_limit",
                            "Текст превышает размер ячейки Excel. Используйте CSV.",
                            422,
                        )
                    cell = WriteOnlyCell(sheet)
                    if field in NUMERIC:
                        cell.value = Decimal(value)
                        cell.number_format = (
                            "0.000"
                            if field.endswith("score")
                            else "#,##0.00"
                            if field == "sum_kzt_internal"
                            else "#,##0"
                        )
                    else:
                        # Explicit strings preserve int64 IDs and prevent formula evaluation.
                        cell.value = value
                        cell.data_type = "s"
                        cell.number_format = "@"
                    cell.font = Font(name="Calibri", size=11, color="283824")
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if row_number % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F2F6EE")
                    cells.append(cell)
                sheet.append(cells)
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{row_number}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
