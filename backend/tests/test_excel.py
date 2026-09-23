import hashlib
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.config import Settings
from backend.app.main import create_app
from backend.app.models import AnalysisRun
from backend.tests.test_analyses import dataset, finished, result_files, results_rows, submit
from backend.tests.test_cases import BIG_GID


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(Settings(storage_dir=tmp_path / "storage"))) as value:
        yield value


def test_excel_preserves_unicode_exact_ids_and_literal_text_without_mutating_sources(client):
    rows = results_rows()
    rows["nodes_roles"][0]["evidence"] = "Вход: 110 000 ₸ · исходящих: 7 переводов"
    rows["nodes_roles"][1]["evidence"] = '=HYPERLINK("https://example.invalid", "test")'
    rows["top_nodes"][0]["why"] = "+1+2"
    run = finished(client, submit(client, dataset(client)["id"], result_files(rows)).json()["id"])
    root = f"/api/v1/analyses/{run['id']}"
    before = {f["name"]: client.get(root + "/exports/" + f["name"]).content for f in run["files"]}
    response = client.get(root + "/excel")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].endswith("spreadsheetml.sheet")
    assert ".xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Приоритеты", "Участники", "Сообщества"]
    sheet = workbook["Участники"]
    assert sheet["A2"].value == str(BIG_GID) and sheet["A2"].data_type == "s"
    assert sheet["A2"].number_format == "@"
    assert sheet["F2"].value == rows["nodes_roles"][0]["evidence"]
    assert sheet["F3"].value.startswith("=HYPERLINK") and sheet["F3"].data_type == "s"
    assert sheet["C2"].value == 0.5 and sheet["C2"].data_type == "n"
    assert sheet.freeze_panes == "B2" and sheet.auto_filter.ref == "A1:F6"
    assert workbook["Приоритеты"]["B2"].value == str(BIG_GID)
    assert workbook["Приоритеты"]["E2"].value == "+1+2"
    assert str(BIG_GID) in workbook["Сообщества"]["E2"].value
    assert sheet.max_row == 6 and workbook["Сообщества"].max_row == 2
    for file in run["files"]:
        after = client.get(root + "/exports/" + file["name"]).content
        assert after == before[file["name"]]
        assert hashlib.sha256(after).hexdigest() == file["sha256"]


def test_excel_rejects_failed_or_tampered_snapshots(client):
    run = finished(client, submit(client, dataset(client)["id"]).json()["id"])
    root = f"/api/v1/analyses/{run['id']}"
    path = client.app.state.analyses.download(run["id"], "clusters.csv")
    path.write_text("tampered", encoding="utf-8")
    response = client.get(root + "/excel")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "export_changed"
    with client.app.state.database.sessions() as db:
        record = db.get(AnalysisRun, run["id"])
        record.status = "failed"
        db.commit()
    assert client.get(root + "/excel").status_code == 409
